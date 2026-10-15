"""
Builds data/processed/dataset.json from raw AQR factor files and
yfinance-derived equity index prices.

Run: python3 scripts/build_data.py
"""
import json
import csv
import datetime as dt
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "processed" / "dataset.json"


def month_key(d):
    return f"{d.year:04d}-{d.month:02d}"


def parse_momentum(path):
    """AQR Momentum Indices, Monthly -> U.S. Large Cap column."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Returns"]
    rows = list(ws.iter_rows(values_only=True))
    out = {}
    for row in rows[2:]:
        date, us_lc = row[0], row[1]
        if date is None or us_lc is None:
            continue
        out[month_key(date)] = float(us_lc)
    return out


def day_key(d):
    return d.strftime("%Y-%m-%d")


def parse_country_factor(path, sheet, country="USA", key=month_key):
    """AQR HML Devil / BAB / QMJ files share the same layout (monthly and daily):
    header row has DATE + ISO3 country codes + aggregates, data starts next row.
    `key` maps a datetime to the dict key (month_key for monthly, day_key for daily)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = None
    header_idx = None
    for i, row in enumerate(rows):
        if row and row[0] == "DATE":
            header = row
            header_idx = i
            break
    col = header.index(country)
    out = {}
    for row in rows[header_idx + 1:]:
        date_val = row[0]
        if not date_val:
            continue
        val = row[col]
        if val is None:
            continue
        if isinstance(date_val, dt.datetime):
            date = date_val
        else:
            date = dt.datetime.strptime(str(date_val), "%m/%d/%Y")
        out[key(date)] = float(val)
    return out


def parse_french_daily(path):
    """Kenneth French daily factors -> (market_total_return, risk_free) keyed by YYYY-MM-DD."""
    mkt, rf = {}, {}
    with open(path) as f:
        for line in f:
            parts = [p.strip() for p in line.strip().split(",")]
            if len(parts) != 5 or not (parts[0].isdigit() and len(parts[0]) == 8):
                continue
            try:
                mkt_rf, rf_val = float(parts[1]), float(parts[4])
            except ValueError:
                continue
            d = f"{parts[0][:4]}-{parts[0][4:6]}-{parts[0][6:]}"
            mkt[d] = (mkt_rf + rf_val) / 100.0
            rf[d] = rf_val / 100.0
    return mkt, rf


def parse_french_factors(path):
    """Kenneth French Data Library, F-F Research Data Factors (monthly).
    Returns (market_total_return, risk_free_rate) dicts keyed by month.
    Market total return = Mkt-RF + RF (CRSP value-weighted market, incl. dividends)."""
    mkt, rf = {}, {}
    with open(path) as f:
        for line in f:
            parts = [p.strip() for p in line.strip().split(",")]
            if len(parts) != 5:
                continue
            datestr = parts[0]
            if not (datestr.isdigit() and len(datestr) == 6):
                continue
            try:
                mkt_rf, rf_val = float(parts[1]), float(parts[4])
            except ValueError:
                continue
            key = f"{datestr[:4]}-{datestr[4:6]}"
            mkt[key] = (mkt_rf + rf_val) / 100.0
            rf[key] = rf_val / 100.0
    return mkt, rf


def parse_price_csv(path):
    """Monthly price level CSV (from yfinance) -> monthly simple return series."""
    dates, prices = [], []
    with open(path) as f:
        reader = csv.reader(f)
        next(reader)  # header
        for row in reader:
            if not row or not row[0] or row[0] == "Date":
                continue
            try:
                d = dt.datetime.strptime(row[0][:10], "%Y-%m-%d")
                p = float(row[1])
            except (ValueError, IndexError):
                continue
            dates.append(d)
            prices.append(p)
    out = {}
    for i in range(1, len(prices)):
        if prices[i - 1] == 0:
            continue
        ret = prices[i] / prices[i - 1] - 1
        out[month_key(dates[i])] = ret
    return out


def parse_price_csv_daily(path, last_date):
    """Daily price CSV (Date, ticker) -> daily simple returns keyed by YYYY-MM-DD,
    dropping any day after last_date (the final month is partial)."""
    dates, prices = [], []
    with open(path) as f:
        reader = csv.reader(f)
        next(reader)
        for row in reader:
            if not row or not row[0] or row[0] == "Date":
                continue
            dates.append(row[0][:10])
            prices.append(float(row[1]))
    return {
        dates[i]: prices[i] / prices[i - 1] - 1
        for i in range(1, len(prices))
        if prices[i - 1] and dates[i] <= last_date
    }


DAILY_SCALE = 1e-5  # daily returns are stored as integers in units of 0.001%


def build_daily_block(daily_series):
    """Compact daily encoding. All series share one calendar (union of trading days,
    stored as day-gaps); each series is {first, values[]} with `first` an index into the
    calendar and values as integer returns (null where the series skips a calendar day)."""
    all_dates = sorted(set().union(*[d.keys() for d in daily_series.values()]))
    base = dt.date.fromisoformat(all_dates[0])
    gaps, prev = [], base
    for d in all_dates:
        cur = dt.date.fromisoformat(d)
        gaps.append((cur - prev).days)
        prev = cur
    idx = {d: i for i, d in enumerate(all_dates)}
    series = {}
    for key, data in daily_series.items():
        first, last = idx[min(data)], idx[max(data)]
        series[key] = {
            "first": first,
            "values": [
                round(data[d] / DAILY_SCALE) if d in data else None
                for d in all_dates[first:last + 1]
            ],
        }
    return {"freq": "d", "scale": DAILY_SCALE, "start": all_dates[0], "gaps": gaps, "series": series}


def stats_line(label, data):
    vals = [v for v in data.values()]
    n = len(vals)
    mean = sum(vals) / n
    sd = (sum((v - mean) ** 2 for v in vals) / (n - 1)) ** 0.5
    keys = sorted(data)
    return f"{label:<10} {keys[0]} -> {keys[-1]}  n={n:>6}  mean={mean * 100:8.4f}%  sd={sd * 100:7.4f}%"


def series_meta(data, label, category, source, note="", is_excess=False):
    keys = sorted(data.keys())
    return {
        "label": label,
        "category": category,
        "source": source,
        "note": note,
        "is_excess": is_excess,
        "start": keys[0] if keys else None,
        "end": keys[-1] if keys else None,
        "data": data,
    }


def main():
    dataset = {
        "generated": dt.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"),
        "series": {},
    }

    mkt_total, rf = parse_french_factors(RAW / "F-F_Research_Data_Factors.csv")
    dataset["series"]["SPX"] = series_meta(
        mkt_total,
        "US Market (CRSP Total Return)",
        "index",
        "Kenneth French Data Library, F-F Research Data Factors (Mkt-RF + RF)",
        "CRSP value-weighted market return including dividends. Used in place of Yahoo's ^GSPC, which is price-only and understates the mean by roughly 2-4%/yr from missing dividend income.",
    )
    dataset["riskfree"] = series_meta(
        rf,
        "Risk-free rate (1-Month T-Bill)",
        "riskfree",
        "Kenneth French Data Library, F-F Research Data Factors (RF column)",
        "1-month T-bill return (Ibbotson through 2024-05, ICE BofA 1-Month T-Bill Index thereafter). Used to convert total returns to excess returns.",
    )
    dataset["series"]["EM"] = series_meta(
        parse_price_csv(RAW / "EM_monthly_price.csv"),
        "Emerging Markets (MSCI EM proxy, since 2003)",
        "index",
        "Yahoo Finance (EEM ETF), dividend-adjusted",
        "iShares MSCI Emerging Markets ETF used as a liquid proxy for the MSCI EM index. Short sample (since 2003): means are estimated very noisily.",
    )
    dataset["series"]["EAFE"] = series_meta(
        parse_price_csv(RAW / "EAFE_monthly_price.csv"),
        "Developed ex-US (MSCI EAFE proxy, since 2001)",
        "index",
        "Yahoo Finance (EFA ETF), dividend-adjusted",
        "iShares MSCI EAFE ETF used as a liquid proxy for the MSCI EAFE index. Short sample (since 2001): means are estimated very noisily.",
    )
    dataset["series"]["MOM"] = series_meta(
        parse_momentum(RAW / "momentum.xlsx"),
        "Momentum (US Large Cap)",
        "factor",
        "AQR Momentum Indices, Monthly",
        "Long-only momentum index return, not a long/short factor.",
    )
    dataset["series"]["VAL"] = series_meta(
        parse_country_factor(RAW / "hml-devil.xlsx", "HML Devil"),
        "Value (HML Devil, US)",
        "factor",
        "AQR: The Devil in HML's Details, Monthly",
        "Long/short self-financing excess return, US equities.",
        is_excess=True,
    )
    dataset["series"]["BAB"] = series_meta(
        parse_country_factor(RAW / "bab.xlsx", "BAB Factors"),
        "Betting Against Beta (US)",
        "factor",
        "AQR: Betting Against Beta, Monthly",
        "Long/short self-financing excess return, US equities.",
        is_excess=True,
    )
    dataset["series"]["QMJ"] = series_meta(
        parse_country_factor(RAW / "qmj.xlsx", "QMJ Factors"),
        "Quality Minus Junk (US)",
        "factor",
        "AQR: Quality Minus Junk, Monthly",
        "Long/short self-financing excess return, US equities.",
        is_excess=True,
    )
    dataset["series"]["BOND"] = series_meta(
        parse_price_csv(RAW / "VBMFX_monthly_price.csv"),
        "US Bonds (VBMFX, proxy for Bloomberg US Aggregate)",
        "bond",
        "Yahoo Finance (VBMFX mutual fund), dividend-adjusted",
        "Vanguard Total Bond Market Index Fund used as a liquid proxy for the Bloomberg US Aggregate Bond Index. Adjusted closes already include reinvested distributions, so monthly changes are total returns; as a fund (not the raw index), these returns are net of a small expense ratio.",
    )

    # ---- daily (feeds only the frequency toggle, tail counts and ACF; monthly stays the default) ----
    DAILY_RAW = RAW / "daily"
    # Drop the partial final month so daily and monthly end on the same convention.
    last_day = f"{max(s['end'] for s in dataset['series'].values())}-31"
    mkt_d, rf_d = parse_french_daily(DAILY_RAW / "F-F_Research_Data_Factors_daily.csv")
    daily = {
        "SPX": mkt_d,
        "RF": rf_d,
        "EM": parse_price_csv_daily(DAILY_RAW / "EEM_daily.csv", last_day),
        "EAFE": parse_price_csv_daily(DAILY_RAW / "EFA_daily.csv", last_day),
        "VAL": parse_country_factor(DAILY_RAW / "hml-devil-daily.xlsx", "HML Devil", key=day_key),
        "BAB": parse_country_factor(DAILY_RAW / "bab-daily.xlsx", "BAB Factors", key=day_key),
        "QMJ": parse_country_factor(DAILY_RAW / "qmj-daily.xlsx", "QMJ Factors", key=day_key),
        "BOND": parse_price_csv_daily(DAILY_RAW / "VBMFX_daily.csv", last_day),
    }
    dataset["daily"] = build_daily_block(daily)

    print("--- monthly ---")
    for key, s in dataset["series"].items():
        print(stats_line(key, s["data"]))
    print(stats_line("RF", rf))
    print("--- daily (MOM has no daily source) ---")
    for key, data in daily.items():
        print(stats_line(key, data))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(dataset, separators=(",", ":")))
    print(f"\nWrote {OUT} ({OUT.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
